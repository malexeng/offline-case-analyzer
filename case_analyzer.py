#!/usr/bin/env python3
"""
==============================================================
  LEGAL CASE ANALYZER - Ontario, Canada (OFFLINE)
  Multi-Modal Evidence Analysis & Report Generator
  Offline LLM via Ollama + OCR-first evidence extraction
==============================================================
"""

import os
import sys
import json
import argparse
import datetime
import re
import hashlib
from pathlib import Path
from typing import Optional, Dict, List, Tuple

import requests

# ---- Required deps ----
MISSING = []
try:
    from docx import Document
    from docx.shared import RGBColor
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
except ImportError:
    MISSING.append("python-docx")

try:
    from PIL import Image
except ImportError:
    MISSING.append("Pillow")

try:
    import fitz  # PyMuPDF
    PDF_SUPPORT = True
except ImportError:
    PDF_SUPPORT = False
    print("⚠  PyMuPDF not found – PDF text extraction disabled (pip install pymupdf)")

try:
    import cv2
    import numpy as np
    VIDEO_SUPPORT = True
except ImportError:
    VIDEO_SUPPORT = False
    print("⚠  OpenCV not found – video analysis disabled (pip install opencv-python)")

# OCR choice:
# - pytesseract requires system tesseract-ocr installed
OCR_SUPPORT = True
try:
    import pytesseract
except ImportError:
    OCR_SUPPORT = False
    print("⚠  pytesseract not found – image OCR disabled (pip install pytesseract). "
          "Also install system 'tesseract-ocr'.")

if MISSING:
    print(f"❌  Missing required packages: {', '.join(MISSING)}")
    print(f"    Run: pip install {' '.join(MISSING)}")
    sys.exit(1)

SUPPORTED_IMAGES = {".jpg", ".jpeg", ".png", ".webp", ".bmp", ".tiff"}
SUPPORTED_DOCS   = {".pdf", ".docx", ".doc", ".txt", ".md", ".rtf"}
SUPPORTED_VIDEO  = {".mp4", ".mov", ".avi", ".mkv", ".webm"}

ONTARIO_LEGAL_FRAMEWORK = """
ONTARIO LEGAL FRAMEWORKS TO CONSIDER (NON-EXHAUSTIVE):

HUMAN RIGHTS / DISCRIMINATION:
- Ontario Human Rights Code (OHRC) – s.1 (services), s.5 (employment)
- HRTO remedies: general damages often in the thousands to tens of thousands; lost wages; public interest remedies

HEALTHCARE-SPECIFIC:
- PHIPA (privacy breaches)
- RHPA / professional misconduct (College processes like CNO/CPSO)
- HCCA (consent)
- Negligence / malpractice (civil)

PRIVACY:
- PHIPA, MFIPPA (context-dependent)

HARASSMENT:
- Criminal Code s.264 (criminal harassment) + civil torts where applicable
"""

SYSTEM_PROMPT = """You are an expert legal analyst and paralegal specializing in Ontario, Canada law.
You analyze evidence (OCR text from screenshots, extracted documents, OCR from video frames)
to identify legal violations, build timelines, identify parties, and assess potential claims and damages.

Rules:
1) FACTUAL: Only state what is supported by evidence text; label uncertainty.
2) STRUCTURED: Follow requested output format exactly (especially JSON).
3) ONTARIO-SPECIFIC: Reference Ontario laws/tribunals where relevant.
4) EVIDENCE-LINKED: Each key claim should reference filenames.
5) TIMELINE: Reconstruct chronology from dates/times in evidence.

If evidence is insufficient, say so and list what would strengthen the case.
"""

# ============================================================
#  Utility
# ============================================================

def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()

def safe_read_text(path: Path) -> str:
    try:
        return path.read_text(encoding="utf-8", errors="replace")
    except Exception as e:
        return f"[Read error: {e}]"

def extract_pdf_text(pdf_path: Path) -> str:
    if not PDF_SUPPORT:
        return f"[PDF text extraction unavailable – install pymupdf. File: {pdf_path.name}]"
    try:
        doc = fitz.open(str(pdf_path))
        parts = []
        for i, page in enumerate(doc, 1):
            t = page.get_text()
            if t.strip():
                parts.append(f"--- Page {i} ---\n{t}")
        doc.close()
        return "\n".join(parts) if parts else "[PDF appears image-based or empty]"
    except Exception as e:
        return f"[PDF extraction error: {e}]"

def extract_docx_text(docx_path: Path) -> str:
    try:
        doc = Document(str(docx_path))
        paras = [p.text for p in doc.paragraphs if p.text.strip()]
        table_rows = []
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    table_rows.append(row_text)
        combined = paras + (["[TABLE DATA:]"] + table_rows if table_rows else [])
        return "\n".join(combined)
    except Exception as e:
        return f"[DOCX extraction error: {e}]"

# ============================================================
#  OCR
# ============================================================

def _preprocess_for_ocr(img_bgr):
    # Convert to grayscale
    gray = cv2.cvtColor(img_bgr, cv2.COLOR_BGR2GRAY)

    # Upscale small images (helps a lot with UI screenshots)
    h, w = gray.shape[:2]
    if max(h, w) < 1400:
        scale = 2
        gray = cv2.resize(gray, (w * scale, h * scale), interpolation=cv2.INTER_CUBIC)

    # Denoise + adaptive threshold
    gray = cv2.bilateralFilter(gray, 9, 75, 75)
    thr = cv2.adaptiveThreshold(
        gray, 255, cv2.ADAPTIVE_THRESH_GAUSSIAN_C,
        cv2.THRESH_BINARY, 31, 8
    )
    return thr

def ocr_image(image_path: Path) -> str:
    if not (OCR_SUPPORT and VIDEO_SUPPORT):
        # VIDEO_SUPPORT implies cv2 available; we use cv2 for preprocess
        return "[OCR unavailable]"
    img = cv2.imread(str(image_path))
    if img is None:
        return ""
    proc = _preprocess_for_ocr(img)
    # psm 6 is decent for blocks of text; tweak if needed
    cfg = "--oem 3 --psm 6"
    text = pytesseract.image_to_string(proc, lang="eng", config=cfg)
    return text.strip()

# ============================================================
#  Video frames -> OCR text
# ============================================================

def extract_video_frames(video_path: Path, max_frames: int = 4) -> List[Tuple[Path, str]]:
    """
    Extract a few frames to temp images (in memory would be nicer, but this is simple).
    Returns list of (temp_path, label).
    """
    if not VIDEO_SUPPORT:
        return []
    cap = cv2.VideoCapture(str(video_path))
    if not cap.isOpened():
        return []
    total = int(cap.get(cv2.CAP_PROP_FRAME_COUNT)) or 0
    fps = cap.get(cv2.CAP_PROP_FPS) or 0.0

    # sample evenly
    if total <= 0:
        cap.release()
        return []

    sample_points = [int(i * total / max_frames) for i in range(max_frames)]
    out = []
    tmp_dir = video_path.parent / ".frames_tmp"
    tmp_dir.mkdir(exist_ok=True)

    for idx, frame_num in enumerate(sample_points):
        cap.set(cv2.CAP_PROP_POS_FRAMES, frame_num)
        ret, frame = cap.read()
        if not ret:
            continue
        ts = (frame_num / fps) if fps > 0 else 0.0

        # shrink a bit to keep OCR fast, but not too small
        h, w = frame.shape[:2]
        max_dim = 1600
        if max(h, w) > max_dim:
            scale = max_dim / max(h, w)
            frame = cv2.resize(frame, (int(w * scale), int(h * scale)), interpolation=cv2.INTER_AREA)

        tmp_path = tmp_dir / f"{video_path.stem}_frame{idx+1}.jpg"
        cv2.imwrite(str(tmp_path), frame, [int(cv2.IMWRITE_JPEG_QUALITY), 85])
        out.append((tmp_path, f"Frame {idx+1} at {ts:.1f}s"))
    cap.release()
    return out

# ============================================================
#  Ollama client
# ============================================================

class OllamaClient:
    def __init__(self, host: str = "http://localhost:11434"):
        self.host = host.rstrip("/")

    def generate(self, model: str, prompt: str, num_ctx: int = 4096, temperature: float = 0.2) -> str:
        payload = {
            "model": model,
            "prompt": prompt,
            "stream": False,
            "options": {
                "num_ctx": num_ctx,
                "temperature": temperature,
                "top_p": 0.9,
                "repeat_penalty": 1.1,
            }
        }
        r = requests.post(f"{self.host}/api/generate", json=payload, timeout=900)
        r.raise_for_status()
        return r.json().get("response", "")

# ============================================================
#  Evidence collector
# ============================================================

class EvidenceCollector:
    def __init__(self, case_folder: str):
        self.case_folder = Path(case_folder)
        if not self.case_folder.exists():
            raise FileNotFoundError(f"Case folder not found: {case_folder}")
        self.evidence = {"images": [], "documents": [], "videos": [], "other": []}

    def scan(self) -> dict:
        print(f"\n📁  Scanning case folder: {self.case_folder}")
        for path in sorted(self.case_folder.rglob("*")):
            if path.is_file():
                suf = path.suffix.lower()
                if suf in SUPPORTED_IMAGES:
                    self.evidence["images"].append(path)
                elif suf in SUPPORTED_DOCS:
                    self.evidence["documents"].append(path)
                elif suf in SUPPORTED_VIDEO:
                    self.evidence["videos"].append(path)
                else:
                    self.evidence["other"].append(path)

        total = sum(len(v) for v in self.evidence.values())
        print(f"   ✅  Found {total} files:")
        print(f"       📸  Images   : {len(self.evidence['images'])}")
        print(f"       📄  Documents: {len(self.evidence['documents'])}")
        print(f"       🎬  Videos   : {len(self.evidence['videos'])}")
        return self.evidence

# ============================================================
#  JSON robustness helpers
# ============================================================

def extract_first_json_object(text: str) -> str:
    text = text.strip()
    # remove fences
    text = re.sub(r"^```json\s*", "", text)
    text = re.sub(r"^```\s*", "", text)
    text = re.sub(r"\s*```$", "", text)

    # greedy: first { ... } block
    m = re.search(r"\{.*\}", text, flags=re.DOTALL)
    return m.group(0) if m else ""

def parse_json_strict(raw: str) -> dict:
    candidate = extract_first_json_object(raw)
    if not candidate:
        raise json.JSONDecodeError("No JSON object found", raw, 0)
    return json.loads(candidate)

# ============================================================
#  Case analyzer (offline)
# ============================================================

class CaseAnalyzerOffline:
    def __init__(self, ollama: OllamaClient, case_type: str, text_model: str, num_ctx: int):
        self.ollama = ollama
        self.case_type = case_type
        self.text_model = text_model
        self.num_ctx = num_ctx

    def _call_llm(self, prompt: str, temperature: float = 0.2) -> str:
        full = SYSTEM_PROMPT + "\n\n" + ONTARIO_LEGAL_FRAMEWORK + "\n\n" + prompt
        return self.ollama.generate(self.text_model, full, num_ctx=self.num_ctx, temperature=temperature)

    def analyze_images(self, image_paths: List[Path]) -> str:
        if not image_paths:
            return "No images provided."

        print(f"\n📸  OCR+Analyzing {len(image_paths)} images...")
        blocks = []
        for p in image_paths:
            ocr = ocr_image(p) if OCR_SUPPORT else "[OCR unavailable]"
            blocks.append(f"""
━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
IMAGE FILE: {p.name}
SHA256: {sha256_file(p)}
[OCR TEXT BEGIN]
{ocr[:12000]}
[OCR TEXT END]
━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
""")

        prompt = f"""
Case type: {self.case_type}

You are given OCR text extracted from image evidence files (screenshots/photos).
For EACH file:
1) Summarize what it contains (1-3 sentences).
2) Extract key entities: names, organizations, emails, phone numbers, addresses.
3) Extract ALL dates/times and what happened at each.
4) Note any statements that indicate refusals, discrimination, privacy issues, misconduct, threats, etc.
5) Assign an evidence strength: Strong/Moderate/Weak (based on clarity/specificity).
6) Output results as a bullet list grouped by filename.

OCR INPUT:
{''.join(blocks)}
"""
        return self._call_llm(prompt, temperature=0.2)

    def analyze_documents(self, doc_paths: List[Path]) -> str:
        if not doc_paths:
            return "No documents provided."

        print(f"\n📄  Extracting+Analyzing {len(doc_paths)} documents...")
        blocks = []
        for p in doc_paths:
            suf = p.suffix.lower()
            if suf == ".pdf":
                text = extract_pdf_text(p)
            elif suf in (".docx", ".doc"):
                text = extract_docx_text(p)
            else:
                text = safe_read_text(p)

            blocks.append(f"""
━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
DOCUMENT FILE: {p.name}
SHA256: {sha256_file(p)}
[TEXT BEGIN]
{text[:16000]}
{"...[TRUNCATED]" if len(text) > 16000 else ""}
[TEXT END]
━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
""")

        prompt = f"""
Case type: {self.case_type}

Analyze these extracted document texts.
1) Parties identified (names/titles/orgs)
2) Key facts and admissions
3) All dates -> timeline entries
4) Potential Ontario legal issues indicated
5) Gaps/missing documents

DOCUMENT INPUT:
{''.join(blocks)}
"""
        return self._call_llm(prompt, temperature=0.2)

    def analyze_videos(self, video_paths: List[Path]) -> str:
        if not video_paths:
            return "No videos provided."
        if not VIDEO_SUPPORT:
            return "Video analysis unavailable (opencv not installed)."

        print(f"\n🎬  Extracting frames + OCR+Analyzing {len(video_paths)} videos...")
        blocks = []
        for v in video_paths:
            frames = extract_video_frames(v, max_frames=4)
            if not frames:
                blocks.append(f"VIDEO: {v.name} — [Could not extract frames]\n")
                continue

            blocks.append(f"\nVIDEO FILE: {v.name}\nSHA256: {sha256_file(v)}\n")
            for frame_path, label in frames:
                ocr = ocr_image(frame_path) if OCR_SUPPORT else "[OCR unavailable]"
                blocks.append(f"""
--- {label} ({frame_path.name}) ---
[OCR TEXT BEGIN]
{ocr[:8000]}
[OCR TEXT END]
""")

        prompt = f"""
Case type: {self.case_type}

You are given OCR text from sampled video frames (not audio).
1) Summarize what the video likely shows based on frame OCR.
2) Extract names/dates/times/places shown.
3) Note any legally relevant actions or communications.
4) Evidence strength and gaps (e.g., audio missing).

VIDEO FRAME OCR INPUT:
{''.join(blocks)}
"""
        return self._call_llm(prompt, temperature=0.2)

    def synthesize_case(self, case_name: str, case_description: str,
                        image_analysis: str, doc_analysis: str, video_analysis: str,
                        victim_name: str) -> dict:
        print(f"\n⚖️   Synthesizing legal conclusions (JSON)...")

        synthesis_prompt = f"""
Produce the FINAL STRUCTURED LEGAL ANALYSIS for Ontario, Canada.

CASE NAME: {case_name}
VICTIM: {victim_name}
CASE DESCRIPTION: {case_description}

EVIDENCE ANALYSIS (OCR + docs):
[IMAGES]
{image_analysis[:12000]}

[DOCUMENTS]
{doc_analysis[:12000]}

[VIDEOS]
{video_analysis[:8000]}

REQUIRED OUTPUT:
Return ONLY valid JSON with these keys and structure:

{{
  "case_summary": "...",
  "parties": {{
    "victim": {{
      "name": "{victim_name}",
      "role": "Complainant / Plaintiff",
      "protected_grounds": ["..."]
    }},
    "respondents": [
      {{
        "name": "...",
        "title": "...",
        "organization": "...",
        "role_in_case": "...",
        "evidence_links": ["file1.png", "doc1.pdf"]
      }}
    ]
  }},
  "timeline": [
    {{
      "date": "YYYY-MM-DD or approximate",
      "event": "...",
      "parties_involved": ["..."],
      "evidence": ["..."],
      "legal_significance": "..."
    }}
  ],
  "violations": [
    {{
      "violation_type": "...",
      "statute": "...",
      "who_violated": "...",
      "what_they_did": "...",
      "when": "...",
      "where": "...",
      "evidence": ["..."],
      "severity": "Low / Medium / High / Critical",
      "notes": "..."
    }}
  ],
  "legal_pathways": [
    {{
      "pathway": "...",
      "body": "...",
      "applicable_law": "...",
      "limitation_period": "...",
      "process": "...",
      "likelihood_of_success": "Strong / Moderate / Weak",
      "reasoning": "..."
    }}
  ],
  "damages": {{
    "general_damages": {{
      "description": "...",
      "estimated_range": "...",
      "basis": "...",
      "supporting_evidence": ["..."]
    }},
    "special_damages": {{
      "description": "...",
      "items": [{{"item":"...","amount":"...","evidence":"..."}}],
      "total_estimated": "..."
    }},
    "lost_income": {{
      "applicable": true,
      "description": "...",
      "estimated": "..."
    }},
    "future_care": {{
      "applicable": true,
      "description": "...",
      "estimated": "..."
    }},
    "aggravated_punitive": {{
      "applicable": true,
      "description": "...",
      "estimated": "..."
    }},
    "total_claim_range": "...",
    "notes": "..."
  }},
  "evidence_inventory": [
    {{
      "filename": "...",
      "type": "...",
      "supports_violations": ["indices like 0,1,2"],
      "strength": "Strong / Moderate / Weak",
      "notes": "..."
    }}
  ],
  "immediate_action_items": ["..."],
  "evidence_gaps": ["..."],
  "professional_referrals": ["..."]
}}

ABSOLUTE RULES:
- JSON only. No extra text.
- If unsure, write "Unknown" or leave conservative ranges, and explain in notes.
"""

        # JSON reliability: try parse; if fails, force model to fix
        attempts = 3
        last_raw = ""
        for i in range(attempts):
            raw = self._call_llm(synthesis_prompt, temperature=0.1)
            last_raw = raw
            try:
                return parse_json_strict(raw)
            except Exception as e:
                synthesis_prompt = f"""
Your last output was invalid JSON. Fix it.

Return ONLY valid JSON (single object), no commentary.

INVALID OUTPUT:
{raw}
"""
        return {"raw_synthesis": last_raw, "parse_error": "Could not parse valid JSON after retries"}

# ============================================================
#  Report Generator (DOCX)
# ============================================================

class ReportGenerator:
    def __init__(self):
        self.doc = Document()
        self._setup_styles()

    def _setup_styles(self):
        from docx.shared import Cm
        for section in self.doc.sections:
            section.top_margin = Cm(2.5)
            section.bottom_margin = Cm(2.5)
            section.left_margin = Cm(3)
            section.right_margin = Cm(2.5)

    def _heading(self, text: str, level: int = 1):
        p = self.doc.add_heading(text, level=level)
        if p.runs:
            if level == 1:
                p.runs[0].font.color.rgb = RGBColor(0x1a, 0x3a, 0x5c)
            elif level == 2:
                p.runs[0].font.color.rgb = RGBColor(0x2e, 0x6d, 0xab)
        return p

    def _para(self, text: str, bold: bool = False, italic: bool = False):
        p = self.doc.add_paragraph(text)
        for run in p.runs:
            run.bold = bold
            run.italic = italic
        return p

    def _table_row(self, table, cells: list, header: bool = False):
        row = table.add_row()
        for cell, text in zip(row.cells, cells):
            cell.text = str(text)
            if header:
                for run in cell.paragraphs[0].runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                shading = OxmlElement('w:shd')
                shading.set(qn('w:fill'), '1a3a5c')
                shading.set(qn('w:val'), 'clear')
                cell._tc.get_or_add_tcPr().append(shading)
        return row

    def generate(self, case_name: str, analysis: dict, output_path: str,
                 image_analysis_raw: str = "", doc_analysis_raw: str = "") -> str:
        now = datetime.datetime.now().strftime("%B %d, %Y")

        self._heading("LEGAL CASE ANALYSIS REPORT", 1)
        self._para(f"Case: {case_name}", bold=True)
        self._para("Jurisdiction: Ontario, Canada")
        self._para(f"Generated: {now}")
        self._para("Prepared by: Offline AI-Assisted Legal Case Analyzer")
        self.doc.add_paragraph()

        p = self.doc.add_paragraph()
        p.add_run("⚠  DISCLAIMER: ").bold = True
        p.add_run(
            "This report is AI-generated for informational purposes only and does NOT constitute legal advice. "
            "Consult a qualified Ontario lawyer before taking action."
        ).italic = True
        self.doc.add_paragraph()

        self._heading("1. CASE SUMMARY", 1)
        self._para(analysis.get("case_summary", "N/A"))
        self.doc.add_paragraph()

        self._heading("2. PARTIES", 1)
        parties = analysis.get("parties", {})
        self._heading("2.1 Victim / Complainant", 2)
        victim = parties.get("victim", {})
        self._para(f"Name: {victim.get('name','N/A')}")
        self._para(f"Role: {victim.get('role','N/A')}")
        grounds = victim.get("protected_grounds", [])
        if grounds:
            self._para(f"Protected Grounds (OHRC): {', '.join(grounds)}")
        self.doc.add_paragraph()

        self._heading("2.2 Respondents / Defendants", 2)
        respondents = parties.get("respondents", [])
        if respondents:
            table = self.doc.add_table(rows=1, cols=5)
            table.style = "Table Grid"
            self._table_row(table, ["Name","Title","Organization","Role in Case","Evidence"], header=True)
            for r in respondents:
                self._table_row(table, [
                    r.get("name",""),
                    r.get("title",""),
                    r.get("organization",""),
                    r.get("role_in_case",""),
                    ", ".join(r.get("evidence_links", []))
                ])
        self.doc.add_paragraph()

        self._heading("3. CHRONOLOGICAL TIMELINE", 1)
        timeline = analysis.get("timeline", [])
        if timeline:
            table = self.doc.add_table(rows=1, cols=4)
            table.style = "Table Grid"
            self._table_row(table, ["Date","Event","Parties Involved","Evidence"], header=True)
            for t in timeline:
                self._table_row(table, [
                    t.get("date",""),
                    t.get("event",""),
                    ", ".join(t.get("parties_involved", [])),
                    ", ".join(t.get("evidence", []))
                ])
        self.doc.add_paragraph()

        self._heading("4. IDENTIFIED LEGAL VIOLATIONS", 1)
        violations = analysis.get("violations", [])
        for i, v in enumerate(violations, 1):
            self._heading(f"Violation {i}: {v.get('violation_type','Unknown')}", 2)
            tbl = self.doc.add_table(rows=0, cols=2)
            tbl.style = "Table Grid"
            rows_data = [
                ("Statute / Law", v.get("statute","")),
                ("Who Violated", v.get("who_violated","")),
                ("Conduct", v.get("what_they_did","")),
                ("When", v.get("when","")),
                ("Where", v.get("where","")),
                ("Severity", v.get("severity","")),
                ("Evidence", ", ".join(v.get("evidence", []))),
                ("Notes", v.get("notes","")),
            ]
            for label, val in rows_data:
                row = tbl.add_row()
                row.cells[0].text = label
                row.cells[0].paragraphs[0].runs[0].bold = True
                row.cells[1].text = str(val)
            self.doc.add_paragraph()

        self._heading("5. LEGAL PATHWAYS & REMEDIES", 1)
        for pw in analysis.get("legal_pathways", []):
            self._heading(f"→ {pw.get('pathway','')}", 2)
            self._para(f"Body: {pw.get('body','')}")
            self._para(f"Applicable Law: {pw.get('applicable_law','')}")
            self._para(f"Limitation Period: {pw.get('limitation_period','')}")
            self._para(f"Likelihood of Success: {pw.get('likelihood_of_success','')}")
            self._para(pw.get("reasoning",""))
            self.doc.add_paragraph()

        self._heading("6. DAMAGES ASSESSMENT", 1)
        damages = analysis.get("damages", {})
        table = self.doc.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        self._table_row(table, ["Category","Estimated Range","Description / Basis"], header=True)

        def get_est(cat, key="estimated"):
            return cat.get(key, "TBD")

        damage_rows = [
            ("General Damages",
             damages.get("general_damages", {}).get("estimated_range","TBD"),
             damages.get("general_damages", {}).get("basis","")),
            ("Special Damages",
             damages.get("special_damages", {}).get("total_estimated","TBD"),
             damages.get("special_damages", {}).get("description","")),
            ("Lost Income",
             damages.get("lost_income", {}).get("estimated","TBD") if damages.get("lost_income", {}).get("applicable") else "N/A",
             damages.get("lost_income", {}).get("description","")),
            ("Future Care",
             damages.get("future_care", {}).get("estimated","TBD") if damages.get("future_care", {}).get("applicable") else "N/A",
             damages.get("future_care", {}).get("description","")),
            ("Aggravated / Punitive",
             damages.get("aggravated_punitive", {}).get("estimated","TBD") if damages.get("aggravated_punitive", {}).get("applicable") else "N/A",
             damages.get("aggravated_punitive", {}).get("description","")),
        ]
        for row in damage_rows:
            self._table_row(table, list(row))

        total = damages.get("total_claim_range","TBD")
        p = self.doc.add_paragraph()
        p.add_run("TOTAL ESTIMATED CLAIM RANGE: ").bold = True
        p.add_run(total).bold = True
        self.doc.add_paragraph()

        self._heading("7. EVIDENCE INVENTORY", 1)
        inv = analysis.get("evidence_inventory", [])
        if inv:
            table = self.doc.add_table(rows=1, cols=4)
            table.style = "Table Grid"
            self._table_row(table, ["File","Type","Strength","What It Proves"], header=True)
            for e in inv:
                self._table_row(table, [e.get("filename",""), e.get("type",""), e.get("strength",""), e.get("notes","")])
        self.doc.add_paragraph()

        self._heading("8. IMMEDIATE ACTION ITEMS", 1)
        for item in analysis.get("immediate_action_items", []):
            self.doc.add_paragraph(str(item), style="List Number")
        self.doc.add_paragraph()

        self._heading("9. EVIDENCE GAPS & RECOMMENDATIONS", 1)
        for gap in analysis.get("evidence_gaps", []):
            self.doc.add_paragraph(str(gap), style="List Bullet")
        self.doc.add_paragraph()

        self._heading("10. PROFESSIONAL REFERRALS", 1)
        for ref in analysis.get("professional_referrals", []):
            self.doc.add_paragraph(str(ref), style="List Bullet")
        self.doc.add_paragraph()

        if image_analysis_raw or doc_analysis_raw:
            self._heading("APPENDIX A – RAW AI EVIDENCE ANALYSIS", 1)
            if image_analysis_raw:
                self._heading("A.1 Image OCR + Analysis", 2)
                self._para(image_analysis_raw[:8000] + ("..." if len(image_analysis_raw) > 8000 else ""))
            if doc_analysis_raw:
                self._heading("A.2 Document Analysis", 2)
                self._para(doc_analysis_raw[:8000] + ("..." if len(doc_analysis_raw) > 8000 else ""))

        self.doc.save(output_path)
        print(f"\n✅  Report saved: {output_path}")
        return output_path

# ============================================================
#  Save JSON
# ============================================================

def save_json_report(analysis: dict, output_path: str):
    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(analysis, f, indent=2, ensure_ascii=False)
    print(f"✅  JSON report saved: {output_path}")

# ============================================================
#  MAIN
# ============================================================

def main():
    parser = argparse.ArgumentParser(description="Legal Case Analyzer (Offline) – Ollama + OCR")
    parser.add_argument("--case-folder", required=True)
    parser.add_argument("--case-name", default="Legal Case Analysis")
    parser.add_argument("--victim", default="The Complainant")
    parser.add_argument("--case-type", default="general",
                        choices=["general", "healthcare", "employment", "harassment", "privacy", "housing"])
    parser.add_argument("--description", default="")
    parser.add_argument("--output-folder", default=None)
    parser.add_argument("--no-report", action="store_true")
    parser.add_argument("--ollama-host", default="http://localhost:11434")
    parser.add_argument("--text-model", default="llama3.2:latest")
    parser.add_argument("--num-ctx", type=int, default=4096)

    args = parser.parse_args()

    print("""
╔══════════════════════════════════════════════════════════╗
║   LEGAL CASE ANALYZER (OFFLINE) – Ontario, Canada        ║
║   OCR-first + Ollama Local LLM                           ║
╚══════════════════════════════════════════════════════════╝
""")

    output_folder = Path(args.output_folder) if args.output_folder else Path(args.case_folder)
    output_folder.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    safe_name = re.sub(r'[^\w\s-]', '', args.case_name).strip().replace(' ', '_')

    collector = EvidenceCollector(args.case_folder)
    evidence = collector.scan()
    if not any(len(v) > 0 for v in evidence.values()):
        print("❌  No supported evidence files found in the case folder.")
        sys.exit(1)

    ollama = OllamaClient(args.ollama_host)
    analyzer = CaseAnalyzerOffline(
        ollama=ollama,
        case_type=args.case_type,
        text_model=args.text_model,
        num_ctx=args.num_ctx
    )

    print("\n" + "═" * 60)
    print(" PHASE 1: EVIDENCE EXTRACTION & ANALYSIS (OFFLINE)")
    print("═" * 60)

    image_analysis = analyzer.analyze_images(evidence["images"])
    doc_analysis   = analyzer.analyze_documents(evidence["documents"])
    video_analysis = analyzer.analyze_videos(evidence["videos"])

    print("\n" + "═" * 60)
    print(" PHASE 2: LEGAL SYNTHESIS (JSON)")
    print("═" * 60)

    final_analysis = analyzer.synthesize_case(
        case_name=args.case_name,
        case_description=args.description,
        image_analysis=image_analysis,
        doc_analysis=doc_analysis,
        video_analysis=video_analysis,
        victim_name=args.victim
    )

    print("\n" + "═" * 60)
    print(" PHASE 3: OUTPUT")
    print("═" * 60)

    json_path = output_folder / f"{safe_name}_{timestamp}_analysis.json"
    save_json_report(final_analysis, str(json_path))

    if not args.no_report and "raw_synthesis" not in final_analysis and "parse_error" not in final_analysis:
        docx_path = output_folder / f"{safe_name}_{timestamp}_report.docx"
        gen = ReportGenerator()
        gen.generate(args.case_name, final_analysis, str(docx_path), image_analysis, doc_analysis)

    print(f"\n📁  Output folder: {output_folder}")
    print("✅  Done.")

if __name__ == "__main__":
    main()
